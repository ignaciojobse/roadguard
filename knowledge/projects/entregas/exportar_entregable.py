"""Exporta un entregable en Markdown a DOCX y PDF con la marca RoadGuard.

Uso:
    python knowledge/projects/entregas/exportar_entregable.py \
        knowledge/projects/entregas/bitacora-1.md

Genera <nombre>.docx (python-docx) y <nombre>.pdf (Chrome headless) junto al
Markdown. La ruta de Chrome se puede cambiar con la variable CHROME. Entiende
el subconjunto de Markdown del knowledge: titulos, parrafos, citas, listas,
tablas, negrita, cursiva, codigo y enlaces (se exporta su texto). Los
comentarios HTML (<!-- -->) son notas internas y no se exportan.
"""

from __future__ import annotations

import base64
import html
import os
import re
import subprocess
import sys
import tempfile
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONDO = "0B0F14"
SUPERFICIE = "E8EDF3"
TEXTO = RGBColor(0x14, 0x1B, 0x24)
SECUNDARIO = RGBColor(0x5A, 0x66, 0x72)
CLARO = RGBColor(0xF2, 0xF5, 0xF8)
GRIS = RGBColor(0x9A, 0xA6, 0xB2)
ACENTO = "4C9AFF"
FUENTE = "Arial"
LOGO = Path(__file__).parent / "marca" / "roadguard-logo-oscuro.png"
PIE = "TEL-328 Procesamiento Digital de Imágenes · Castillo · Jobse · Ricci · Zúñiga"
CHROME = os.environ.get(
    "CHROME", "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"
)
INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


@dataclass
class Bloque:
    tipo: str  # titulo, h2, h3, parrafo, cita, vineta, numero, tabla
    texto: str = ""
    filas: list[list[str]] | None = None


def leer_bloques(ruta_md: Path) -> list[Bloque]:
    lineas = ruta_md.read_text(encoding="utf-8").splitlines()
    bloques: list[Bloque] = []
    i = 0
    while i < len(lineas):
        linea = lineas[i]
        if linea.strip().startswith("<!--"):
            while i < len(lineas) and "-->" not in lineas[i]:
                i += 1
            i += 1
            continue
        if linea.startswith("|"):
            filas = []
            while i < len(lineas) and lineas[i].startswith("|"):
                if not re.fullmatch(r"\|[\s:|-]+\|?", lineas[i].strip()):
                    filas.append([c.strip() for c in lineas[i].strip().strip("|").split("|")])
                i += 1
            bloques.append(Bloque("tabla", filas=filas))
            continue
        if linea.startswith("# "):
            bloques.append(Bloque("titulo", linea[2:].strip()))
        elif linea.startswith("### "):
            bloques.append(Bloque("h3", linea[4:].strip()))
        elif linea.startswith("## "):
            bloques.append(Bloque("h2", linea[3:].strip()))
        elif linea.startswith(">"):
            bloques.append(Bloque("cita", linea.lstrip(">").strip()))
        elif re.match(r"^\s*\d+\. ", linea):
            bloques.append(Bloque("numero", re.sub(r"^\s*\d+\. ", "", linea)))
        elif re.match(r"^\s*- ", linea):
            bloques.append(Bloque("vineta", re.sub(r"^\s*- ", "", linea)))
        elif linea.strip():
            bloques.append(Bloque("parrafo", linea.strip()))
        i += 1
    return bloques


def _trozos(texto: str):
    """(texto, negrita, cursiva, codigo) por trozo de formato en linea."""
    for trozo in INLINE.split(texto):
        if not trozo:
            continue
        if trozo.startswith("**") and trozo.endswith("**"):
            yield trozo[2:-2].replace("`", ""), True, False, False
        elif trozo.startswith("`") and trozo.endswith("`"):
            yield trozo[1:-1], False, False, True
        elif trozo.startswith("[") and "](" in trozo:
            yield trozo[1:trozo.index("](")], False, False, False
        elif trozo.startswith("*") and trozo.endswith("*") and len(trozo) > 2:
            yield trozo[1:-1], False, True, False
        else:
            yield trozo, False, False, False


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------

def _sombrear(celda, color: str) -> None:
    sombra = OxmlElement("w:shd")
    for k, v in (("val", "clear"), ("color", "auto"), ("fill", color)):
        sombra.set(qn(f"w:{k}"), v)
    celda._tc.get_or_add_tcPr().append(sombra)


def _borde_inferior(parrafo, color: str) -> None:
    bordes = OxmlElement("w:pBdr")
    abajo = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "12"), ("space", "4"), ("color", color)):
        abajo.set(qn(f"w:{k}"), v)
    bordes.append(abajo)
    parrafo._p.get_or_add_pPr().append(bordes)


def _runs(parrafo, texto: str, color: RGBColor = TEXTO, tam: float | None = None,
          negrita: bool = False) -> None:
    for trozo, b, it, code in _trozos(texto):
        run = parrafo.add_run(trozo)
        run.bold, run.italic = b or negrita, it
        run.font.name = "Courier New" if code else FUENTE
        run.font.color.rgb = color
        if tam:
            run.font.size = Pt(tam)


def escribir_docx(bloques: list[Bloque], ruta: Path) -> None:
    doc = Document()
    seccion = doc.sections[0]
    seccion.left_margin = seccion.right_margin = Cm(2.2)
    doc.styles["Normal"].font.name = FUENTE
    doc.styles["Normal"].font.size = Pt(10.5)
    for b in bloques:
        if b.tipo == "titulo":
            celda = doc.add_table(rows=1, cols=1).cell(0, 0)
            _sombrear(celda, FONDO)
            celda.paragraphs[0].add_run().add_picture(str(LOGO), height=Cm(1.3))
            _runs(celda.add_paragraph(), b.texto, CLARO, 18, negrita=True)
            _runs(celda.add_paragraph(), PIE, GRIS, 9)
            doc.add_paragraph()
        elif b.tipo in ("h2", "h3"):
            h = doc.add_heading(level=1 if b.tipo == "h2" else 2)
            _runs(h, b.texto, TEXTO, 14 if b.tipo == "h2" else 12, negrita=True)
            if b.tipo == "h2":
                _borde_inferior(h, ACENTO)
        elif b.tipo == "tabla":
            tabla = doc.add_table(rows=len(b.filas), cols=len(b.filas[0]))
            tabla.style = "Table Grid"
            tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
            for i, fila in enumerate(b.filas):
                for j, valor in enumerate(fila):
                    celda = tabla.cell(i, j)
                    if i == 0:
                        _sombrear(celda, SUPERFICIE)
                    _runs(celda.paragraphs[0], valor, tam=9, negrita=(i == 0))
            doc.add_paragraph()
        elif b.tipo == "cita":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            _runs(p, b.texto, SECUNDARIO, 9.5)
        elif b.tipo in ("vineta", "numero"):
            estilo = "List Bullet" if b.tipo == "vineta" else "List Number"
            _runs(doc.add_paragraph(style=estilo), b.texto)
        else:
            _runs(doc.add_paragraph(), b.texto)
    doc.save(ruta)


# ---------------------------------------------------------------------------
# PDF via HTML
# ---------------------------------------------------------------------------

def _html_inline(texto: str) -> str:
    partes = []
    for trozo, b, it, code in _trozos(texto):
        t = html.escape(trozo)
        if code:
            t = f"<code>{t}</code>"
        if b:
            t = f"<strong>{t}</strong>"
        if it:
            t = f"<em>{t}</em>"
        partes.append(t)
    return "".join(partes)


CSS = f"""
@page {{ size: A4; margin: 18mm 20mm; }}
body {{ font-family: 'Space Grotesk', Arial, sans-serif; color: #141B24; font-size: 10.5pt; line-height: 1.45; }}
.cabecera {{ background: #{FONDO}; color: #F2F5F8; padding: 18px 22px; border-radius: 6px; margin-bottom: 18px; }}
.cabecera img {{ height: 34px; display: block; margin-bottom: 10px; }}
.cabecera h1 {{ margin: 0; font-size: 20pt; }}
.cabecera p {{ margin: 4px 0 0; color: #9AA6B2; font-size: 9pt; }}
h2 {{ font-size: 14pt; border-bottom: 2px solid #{ACENTO}; padding-bottom: 3px; margin-top: 20px; }}
h3 {{ font-size: 12pt; }}
blockquote {{ color: #5A6672; font-size: 9.5pt; margin: 0 0 10px; padding-left: 10px; border-left: 3px solid #{SUPERFICIE}; }}
table {{ border-collapse: collapse; width: 100%; font-size: 9pt; margin: 6px 0 12px; page-break-inside: auto; }}
th, td {{ border: 1px solid #C9D2DC; padding: 5px 6px; vertical-align: top; text-align: left; }}
th {{ background: #{SUPERFICIE}; }}
tr {{ page-break-inside: avoid; }}
code {{ font-family: ui-monospace, Menlo, monospace; font-size: 9pt; }}
"""


def escribir_html(bloques: list[Bloque]) -> str:
    logo = base64.b64encode(LOGO.read_bytes()).decode()
    cuerpo: list[str] = []
    lista: str | None = None
    for b in bloques:
        etiqueta = {"vineta": "ul", "numero": "ol"}.get(b.tipo)
        if lista and etiqueta != lista:
            cuerpo.append(f"</{lista}>")
            lista = None
        if etiqueta and lista is None:
            cuerpo.append(f"<{etiqueta}>")
            lista = etiqueta
        if b.tipo == "titulo":
            cuerpo.append(
                f'<div class="cabecera"><img src="data:image/png;base64,{logo}">'
                f"<h1>{_html_inline(b.texto)}</h1><p>{html.escape(PIE)}</p></div>"
            )
        elif b.tipo in ("h2", "h3"):
            cuerpo.append(f"<{b.tipo}>{_html_inline(b.texto)}</{b.tipo}>")
        elif b.tipo == "tabla":
            cab = "".join(f"<th>{_html_inline(c)}</th>" for c in b.filas[0])
            filas = "".join(
                "<tr>" + "".join(f"<td>{_html_inline(c)}</td>" for c in f) + "</tr>"
                for f in b.filas[1:]
            )
            cuerpo.append(f"<table><thead><tr>{cab}</tr></thead><tbody>{filas}</tbody></table>")
        elif b.tipo == "cita":
            cuerpo.append(f"<blockquote>{_html_inline(b.texto)}</blockquote>")
        elif etiqueta:
            cuerpo.append(f"<li>{_html_inline(b.texto)}</li>")
        else:
            cuerpo.append(f"<p>{_html_inline(b.texto)}</p>")
    if lista:
        cuerpo.append(f"</{lista}>")
    return (
        '<!doctype html><html lang="es"><head><meta charset="utf-8">'
        f"<style>{CSS}</style></head><body>{''.join(cuerpo)}</body></html>"
    )


def escribir_pdf(bloques: list[Bloque], ruta: Path) -> None:
    with tempfile.TemporaryDirectory(prefix="rg_export_") as tmp:
        pagina = Path(tmp) / "entregable.html"
        pagina.write_text(escribir_html(bloques), encoding="utf-8")
        subprocess.run(
            [CHROME, "--headless=new", "--disable-gpu", "--no-pdf-header-footer",
             f"--print-to-pdf={ruta.resolve()}", pagina.as_uri()],
            check=True, capture_output=True, timeout=120,
        )


def exportar(ruta_md: Path) -> tuple[Path, Path]:
    bloques = leer_bloques(ruta_md)
    ruta_docx, ruta_pdf = ruta_md.with_suffix(".docx"), ruta_md.with_suffix(".pdf")
    escribir_docx(bloques, ruta_docx)
    escribir_pdf(bloques, ruta_pdf)
    return ruta_docx, ruta_pdf


if __name__ == "__main__":
    for salida in exportar(Path(sys.argv[1])):
        print(salida)
